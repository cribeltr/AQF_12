#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Importa la planilla del Programa de Mantenciones Preventivas (.xlsm) a una
base de datos SQLite con las columnas:

    ID, N° Carpeta, N° Inventario, Equipo, Servicio, Unidad, Ubicación,
    Procedencia, Marca, Modelo, Serie, Año Instalación, Vida Útil Residual,
    Clasificación, ENU / Baja

Puntos clave
------------
* El ID corresponde al ORDEN dentro de la planilla.
* Cada equipo se identifica de forma única por su `serie` o su `n_inventario`
  (índices UNIQUE en el esquema).
* Los números de serie e inventario se guardan como TEXTO y se conservan
  TAL CUAL están registrados, respetando los ceros a la izquierda
  ("0024" nunca se convierte en "24").

Uso
---
    python scripts/importar_excel.py \
        --excel data/Programacion_MP_2026.xlsm \
        --db    equipos.db

Si se ejecuta sin argumentos usa esas mismas rutas por defecto.
"""

from __future__ import annotations

import argparse
import sqlite3
import sys
import unicodedata
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("Falta la dependencia 'openpyxl'. Instálala con: pip install -r requirements.txt")


# --------------------------------------------------------------------------- #
# Configuración
# --------------------------------------------------------------------------- #

RAIZ = Path(__file__).resolve().parent.parent
EXCEL_POR_DEFECTO = RAIZ / "data" / "Programacion_MP_2026.xlsm"
DB_POR_DEFECTO = RAIZ / "equipos.db"
SCHEMA = RAIZ / "schema.sql"
HOJA_POR_DEFECTO = "PMP_2026"

# Encabezado de la planilla (texto normalizado) -> columna física en SQLite.
# La normalización quita acentos, pasa a minúsculas y colapsa espacios, de
# modo que "Año Instalación", "AÑO INSTALACION " o "año  instalación" mapeen
# todas a la misma columna.
MAPA_COLUMNAS = {
    "id": "id",
    "n carpeta": "n_carpeta",
    "n inventario": "n_inventario",
    "equipo": "equipo",
    "servicio": "servicio",
    "unidad": "unidad",
    "ubicacion": "ubicacion",
    "procedencia": "procedencia",
    "marca": "marca",
    "modelo": "modelo",
    "serie": "serie",
    "ano instalacion": "anio_instalacion",
    "vida util residual": "vida_util_residual",
    "clasificacion": "clasificacion",
    "enu / baja": "enu_baja",
    "observacion": "observaciones",   # columna "Observación" de la planilla
}

# Columnas que SIEMPRE se tratan como texto (preservan ceros a la izquierda).
COLUMNAS_TEXTO_ESTRICTO = {"n_inventario", "serie"}

# Valores que en realidad significan "sin dato" y se normalizan a NULL.
PLACEHOLDERS_VACIOS = {"", "n/a", "na", "s/n", "sin dato", "-", "--", "."}

# Columnas que se leen desde la planilla.
ORDEN_COLUMNAS = [
    "id", "n_carpeta", "n_inventario", "equipo", "servicio", "unidad",
    "ubicacion", "procedencia", "marca", "modelo", "serie",
    "anio_instalacion", "vida_util_residual", "clasificacion", "enu_baja",
    "observaciones",
]

# Columnas de la tabla que NO vienen de la planilla y se inicializan vacías
# (las completa el usuario desde la interfaz / aplicar_notas.py).
COLUMNAS_EXTRA = ["notas", "registros", "pendiente", "ultima_intervencion", "proximo_vencimiento"]
COLUMNAS_DB = ORDEN_COLUMNAS + COLUMNAS_EXTRA


# --------------------------------------------------------------------------- #
# Utilidades
# --------------------------------------------------------------------------- #

def normalizar_encabezado(texto: object) -> str:
    """Normaliza un encabezado: sin acentos, minúsculas, espacios colapsados."""
    if texto is None:
        return ""
    s = str(texto)
    s = "".join(c for c in unicodedata.normalize("NFD", s)
                if unicodedata.category(c) != "Mn")  # quita acentos
    s = s.lower().replace("°", "").replace("º", "")
    return " ".join(s.split())


def a_texto_fiel(valor: object) -> str | None:
    """
    Convierte un valor de celda a texto SIN perder información.

    * str  -> se devuelve tal cual (conserva ceros a la izquierda: "0024").
    * int  -> str(int)
    * float entero (13042.0) -> "13042"  (evita el sufijo ".0")
    * float real -> str(float)
    """
    if valor is None:
        return None
    if isinstance(valor, str):
        s = valor.strip()
        return s or None
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor)


def limpiar(valor: object) -> object:
    """Pasa los placeholders vacíos a None; recorta espacios de los strings."""
    if isinstance(valor, str):
        valor = valor.strip()
        if valor.lower() in PLACEHOLDERS_VACIOS:
            return None
        return valor or None
    return valor


def encontrar_fila_encabezado(ws, max_filas: int = 15) -> int:
    """Devuelve el número de fila (1-based) que contiene los encabezados."""
    mejor_fila, mejor_aciertos = None, 0
    for fila in range(1, max_filas + 1):
        aciertos = 0
        for celda in ws[fila]:
            if normalizar_encabezado(celda.value) in MAPA_COLUMNAS:
                aciertos += 1
        if aciertos > mejor_aciertos:
            mejor_fila, mejor_aciertos = fila, aciertos
    if mejor_fila is None or mejor_aciertos < 5:
        raise RuntimeError("No se encontró la fila de encabezados en la planilla.")
    return mejor_fila


def mapear_columnas(ws, fila_encabezado: int) -> dict[str, int]:
    """Mapea columna_sqlite -> índice de columna (1-based) en la planilla."""
    mapa: dict[str, int] = {}
    for celda in ws[fila_encabezado]:
        clave = normalizar_encabezado(celda.value)
        if clave in MAPA_COLUMNAS:
            destino = MAPA_COLUMNAS[clave]
            mapa.setdefault(destino, celda.column)  # primera aparición gana
    faltantes = [c for c in ORDEN_COLUMNAS if c not in mapa]
    if faltantes:
        raise RuntimeError(f"Faltan columnas en la planilla: {faltantes}")
    return mapa


# --------------------------------------------------------------------------- #
# Lógica principal
# --------------------------------------------------------------------------- #

def leer_filas(excel: Path, hoja: str) -> list[dict]:
    """Lee la planilla y devuelve una lista de dicts listos para insertar."""
    # data_only=True -> usa los valores calculados (p. ej. "Vida Útil Residual"
    # es una fórmula en Excel y necesitamos su resultado, no el texto =IF(...)).
    wb = openpyxl.load_workbook(excel, data_only=True, read_only=True)
    if hoja not in wb.sheetnames:
        raise RuntimeError(f"La hoja '{hoja}' no existe. Hojas: {wb.sheetnames}")
    ws = wb[hoja]

    fila_encabezado = encontrar_fila_encabezado(ws)
    mapa = mapear_columnas(ws, fila_encabezado)

    registros: list[dict] = []
    contador_orden = 0
    for valores in ws.iter_rows(min_row=fila_encabezado + 1, values_only=True):
        # Indexar por número de columna (1-based) para tomar solo las mapeadas.
        celdas = {i: v for i, v in enumerate(valores, start=1)}

        crudo = {col: celdas.get(idx) for col, idx in mapa.items()}

        # Saltar filas totalmente vacías (sin equipo, serie ni inventario).
        if not any(a_texto_fiel(crudo.get(c)) for c in ("equipo", "serie", "n_inventario")):
            continue

        contador_orden += 1
        registro: dict[str, object] = {}
        for col in ORDEN_COLUMNAS:
            valor = crudo.get(col)
            if col in COLUMNAS_TEXTO_ESTRICTO:
                # Texto fiel + limpieza de placeholders (NUNCA a número).
                texto = a_texto_fiel(valor)
                registro[col] = limpiar(texto)
            elif col in ("id", "n_carpeta", "anio_instalacion"):
                registro[col] = valor if isinstance(valor, int) else (
                    int(valor) if isinstance(valor, float) and valor.is_integer() else limpiar(valor)
                )
            else:
                registro[col] = limpiar(a_texto_fiel(valor) if not isinstance(valor, str) else valor)

        # El ID = orden en la planilla. Si la celda ID viene vacía, usamos el
        # contador secuencial.
        if not isinstance(registro["id"], int):
            registro["id"] = contador_orden

        # Columnas que no vienen de la planilla: empiezan vacías.
        for extra in COLUMNAS_EXTRA:
            registro[extra] = None

        registros.append(registro)

    return registros


def crear_base(db: Path) -> sqlite3.Connection:
    """Crea la base aplicando schema.sql (rehace tabla y vista desde cero)."""
    if not SCHEMA.exists():
        raise RuntimeError(f"No se encontró el esquema: {SCHEMA}")
    con = sqlite3.connect(db)
    con.executescript(SCHEMA.read_text(encoding="utf-8"))
    return con


def insertar(con: sqlite3.Connection, registros: list[dict]) -> tuple[int, list[str]]:
    """Inserta los registros y devuelve (insertados, lista_de_conflictos)."""
    columnas = ", ".join(COLUMNAS_DB)
    marcadores = ", ".join(["?"] * len(COLUMNAS_DB))
    sql = f"INSERT INTO equipos ({columnas}) VALUES ({marcadores})"

    insertados, conflictos = 0, []
    for reg in registros:
        try:
            con.execute(sql, [reg[c] for c in COLUMNAS_DB])
            insertados += 1
        except sqlite3.IntegrityError as exc:
            conflictos.append(
                f"ID {reg['id']} (serie={reg['serie']!r}, "
                f"inventario={reg['n_inventario']!r}): {exc}"
            )
    con.commit()
    return insertados, conflictos


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--excel", type=Path, default=EXCEL_POR_DEFECTO,
                        help=f"Ruta del .xlsm (def: {EXCEL_POR_DEFECTO})")
    parser.add_argument("--db", type=Path, default=DB_POR_DEFECTO,
                        help=f"Ruta de la base SQLite (def: {DB_POR_DEFECTO})")
    parser.add_argument("--hoja", default=HOJA_POR_DEFECTO,
                        help=f"Hoja a importar (def: {HOJA_POR_DEFECTO})")
    args = parser.parse_args()

    if not args.excel.exists():
        return print(f"ERROR: no existe el archivo {args.excel}") or 1

    print(f"Leyendo planilla : {args.excel}  (hoja: {args.hoja})")
    registros = leer_filas(args.excel, args.hoja)
    print(f"Filas de equipos : {len(registros)}")

    print(f"Creando base     : {args.db}")
    con = crear_base(args.db)
    insertados, conflictos = insertar(con, registros)

    # Estadísticas rápidas.
    con_serie = sum(1 for r in registros if r["serie"])
    con_inv = sum(1 for r in registros if r["n_inventario"])
    sin_id = sum(1 for r in registros if not r["serie"] and not r["n_inventario"])

    print("-" * 60)
    print(f"Insertados              : {insertados}")
    print(f"Con N° de serie         : {con_serie}")
    print(f"Con N° de inventario    : {con_inv}")
    print(f"Sin serie ni inventario : {sin_id}")
    if conflictos:
        print(f"\nConflictos de unicidad ({len(conflictos)}):")
        for c in conflictos:
            print("  - " + c)
    con.close()
    print("-" * 60)
    print(f"Listo. Base de datos creada en: {args.db}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
